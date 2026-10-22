#!/usr/bin/env python3
"""Regenerate the 10 reference-document fixtures for the LLM correction checks (test_llm.py).

Run with the system Python that has python-docx and LibreOffice installed:

    /usr/bin/python3 -I -B live/stress/fixtures/llm/make_fixtures.py

Output (this directory, committed):
  ref-01..05 *.docx   written directly with python-docx
  ref-06..10 *.pdf    written as temporary .docx sources (not kept) and converted with LibreOffice
                      (soffice --headless --convert-to pdf, private profile /tmp/nslive-lo-llm, 60 s timeout)

Each file is small (roughly 1 to 3 KB of text) and mixes Traditional Chinese with English. The reference
terms the mock corrector and the review checks rely on are: 《吶喊》自序, 魯迅, Gettysburg Address,
Abraham Lincoln, LibriVox.org.
"""

from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
PROFILE = "file:///tmp/nslive-lo-llm"
CJK_FONT = "Noto Serif CJK TC"

FILLER_ZH = (
    "本文件僅供噪音抑制與語音轉寫流程的測試使用，內容為一般性的背景說明，並不代表任何正式的學術立場。"
    "在長篇錄音中，講者會換段、停頓與重述，因此校對人員需要同時參考時間戳記與上下文，才能判斷哪些句子需要修正。"
    "常見的錯誤包括同音字誤判、專有名詞拼寫不一致，以及把網址或書名說成口語形式。"
)
FILLER_EN = (
    "This note is background material for the transcription test suite. Speakers pause, repeat themselves, and "
    "change topics in long recordings, so a reviewer needs the timestamps and the surrounding lines before deciding "
    "whether a suggested change is right. Typical mistakes include homophones, inconsistent spelling of names, and "
    "web addresses read aloud as words."
)


def set_cjk(run: object) -> None:
    """Give the run an East Asian font so LibreOffice renders Traditional Chinese with a CJK face."""
    rpr = run._element.get_or_add_rPr()  # type: ignore[attr-defined]
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        from docx.oxml import OxmlElement
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), CJK_FONT)


def make_docx(path: Path, title: str, paragraphs: list[str]) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    style.font.name = "Liberation Serif"
    head = doc.add_heading(title, level=1)
    for run in head.runs:
        set_cjk(run)
    for text in paragraphs:
        para = doc.add_paragraph()
        run = para.add_run(text)
        set_cjk(run)
    doc.core_properties.title = title
    doc.core_properties.author = "noScribe Live test fixture"
    doc.save(str(path))


DOCX_SOURCES: list[tuple[str, str, list[str]]] = [
    ("ref-01-nahan-preface.docx", "《吶喊》自序（參考摘要）", [
        "《吶喊》自序是魯迅為短篇小說集《吶喊》撰寫的序文，收於一九二三年出版的集子之中。"
        "序中回憶了作者早年棄醫從文的經過，以及他決定以文字喚醒國人的想法。",
        "The preface to Nahan (Call to Arms) by Lu Xun describes why the author turned from medicine to literature.",
        FILLER_ZH,
        "書名請以《吶喊》自序為準；校對時不要寫成近似的字形。Reference title: Nahan preface, Lu Xun.",
        FILLER_EN,
    ]),
    ("ref-02-gettysburg-address.docx", "Gettysburg Address by Abraham Lincoln", [
        "Four score and seven years ago our fathers brought forth on this continent, a new nation, conceived in "
        "Liberty, and dedicated to the proposition that all men are created equal.",
        "Now we are engaged in a great civil war, testing whether that nation, or any nation so conceived and so "
        "dedicated, can long endure. We are met on a great battle-field of that war.",
        "It is for us the living, rather, to be dedicated here to the unfinished work which they who fought here "
        "have thus far so nobly advanced. It is rather for us to be here dedicated to the great task remaining "
        "before us.",
        "这篇演说由 Abraham Lincoln 於一八六三年在賓夕法尼亞州的蓋茨堡發表，常被稱為 Gettysburg Address。",
        FILLER_EN,
    ]),
    ("ref-03-librivox-guide.docx", "LibriVox 朗讀指南", [
        "LibriVox.org 是由志願者錄製公共領域有聲書的計畫，所有錄音都以公共領域授權釋出。",
        "Volunteers record chapters of public domain books; the catalogue lives at LibriVox.org and each reader "
        "uploads files for review before publication.",
        "錄音前請先確認書名、章節與作者的寫法，並在開頭朗讀 LibriVox.org 這個網址，讓聽眾知道資料來源。",
        FILLER_ZH,
        FILLER_EN,
    ]),
    ("ref-04-glossary.docx", "術語表（校對用）", [
        "魯迅：中國現代文學的重要作家，代表作包括《吶喊》與《彷徨》。",
        "吶喊：魯迅的短篇小說集，其自序常被用來理解作者的寫作動機。",
        "狂人日記、阿Q正傳：收錄於《吶喊》中的代表性作品。",
        "Abraham Lincoln: 16th President of the United States. Gettysburg Address by Abraham Lincoln is his famous "
        "speech of November 19, 1863.",
        "LibriVox.org: the volunteer audiobook site; spell the domain exactly as written here.",
    ]),
    ("ref-05-filler-notes.docx", "一般背景筆記", [FILLER_ZH, FILLER_EN, FILLER_ZH, FILLER_EN, FILLER_ZH]),
]

# Sources for the PDFs: converted by LibreOffice, the .docx sources are not kept.
PDF_SOURCES: list[tuple[str, str, list[str]]] = [
    ("ref-06-nahan-notes.pdf", "《吶喊》讀書筆記", [
        "讀《吶喊》自序時，可以注意魯迅如何描寫鐵屋子裡的呼喊，以及他對希望的看法。",
        "The notes below compare the preface with the stories; Lu Xun, Nahan, and the preface should be spelled as "
        "in the reference glossary.",
        FILLER_ZH,
        FILLER_EN,
    ]),
    ("ref-07-lincoln-speech.pdf", "Gettysburg Address notes", [
        "Gettysburg Address by Abraham Lincoln, delivered in 1863, is a short speech of about two hundred seventy "
        "words.",
        "Reading aloud: pause after each sentence, and keep the names Abraham Lincoln and Gettysburg accurate.",
        FILLER_EN,
        FILLER_ZH,
    ]),
    ("ref-08-librivox-faq.pdf", "LibriVox FAQ", [
        "Question: Where are the recordings published? Answer: at LibriVox.org, under public domain terms.",
        "問：錄音要在哪裡發布？答：LibriVox.org。請不要只寫網站名稱，應寫出完整網址。",
        FILLER_EN,
        FILLER_ZH,
    ]),
    ("ref-09-reading-list.pdf", "閱讀書單", [
        "魯迅的《吶喊》《彷徨》與《故事新編》是常見的入門書目，也適合搭配錄音閱讀。",
        "Recommended pairing: the Gettysburg Address by Abraham Lincoln with a short essay on civic speech.",
        FILLER_ZH,
        FILLER_EN,
    ]),
    ("ref-10-filler-notes.pdf", "會議記錄背景", [FILLER_EN, FILLER_ZH, FILLER_EN, FILLER_ZH, FILLER_EN]),
]


def convert_to_pdf(sources: list[Path], outdir: Path) -> None:
    cmd = ["soffice", f"-env:UserInstallation={PROFILE}", "--headless", "--convert-to", "pdf",
           "--outdir", str(outdir), *[str(s) for s in sources]]
    subprocess.run(cmd, check=True, timeout=60, stdout=subprocess.PIPE, stderr=subprocess.PIPE)


def main() -> int:
    for name, title, paragraphs in DOCX_SOURCES:
        make_docx(HERE / name, title, paragraphs)
        print(f"wrote {name}")
    tmp = Path(tempfile.mkdtemp(prefix="nslive-llm-fixtures-"))
    try:
        pdf_inputs = []
        for name, title, paragraphs in PDF_SOURCES:
            src = tmp / name.replace(".pdf", ".docx")
            make_docx(src, title, paragraphs)
            pdf_inputs.append(src)
        convert_to_pdf(pdf_inputs, HERE)
        for name, _title, _paragraphs in PDF_SOURCES:
            if not (HERE / name).is_file():
                raise RuntimeError(f"LibreOffice did not produce {name}")
            print(f"converted {name}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
